from __future__ import annotations

import json
import io
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PR = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def natural_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", value)]


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.iter(W + "p"):
        text = "".join(node.text or "" for node in paragraph.iter(W + "t")).strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def xlsx_text(path: Path) -> dict[str, list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    result: dict[str, list[dict[str, object]]] = {}
    for sheet in workbook.worksheets:
        rows = []
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            values = [cell.value for cell in row]
            if any(value not in (None, "") for value in values):
                rows.append({"row": row_number, "values": values})
        result[sheet.title] = rows
    return result


def pdf_text(path: Path) -> tuple[str, int]:
    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(f"\n\n===== PAGE {index} =====\n{page.extract_text() or ''}")
    return "".join(pages), len(reader.pages)


def pdf_image_manifest(path: Path) -> list[dict[str, object]]:
    reader = PdfReader(path)
    result = []
    for page_number, page in enumerate(reader.pages, start=1):
        images = []
        for index, item in enumerate(page.images):
            width = height = None
            try:
                with Image.open(io.BytesIO(item.data)) as image:
                    width, height = image.size
            except Exception:
                pass
            images.append(
                {
                    "index": index,
                    "name": item.name,
                    "bytes": len(item.data),
                    "width": width,
                    "height": height,
                }
            )
        result.append({"page": page_number, "images": images})
    return result


def extract_pdf_images(path: Path, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(path)
    for page_number, page in enumerate(reader.pages, start=1):
        for index, item in enumerate(page.images):
            suffix = Path(item.name).suffix.lower() or ".bin"
            (output / f"page-{page_number:02d}-{index}{suffix}").write_bytes(item.data)


def pptx_manifest(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        names = {entry.filename for entry in entries}
        slide_names = sorted(
            (
                name
                for name in names
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=natural_key,
        )
        slides = []
        for slide_name in slide_names:
            root = ET.fromstring(archive.read(slide_name))
            shapes = []
            for shape in root.iter(P + "sp"):
                paragraphs = []
                for paragraph in shape.iter(A + "p"):
                    value = "".join(
                        node.text or "" for node in paragraph.iter(A + "t")
                    ).strip()
                    if value:
                        paragraphs.append(value)
                if paragraphs:
                    shapes.append("\n".join(paragraphs))
            text = "\n".join(
                value
                for value in (node.text or "" for node in root.iter(A + "t"))
                if value.strip()
            )
            rel_name = slide_name.replace("slides/", "slides/_rels/") + ".rels"
            media = []
            if rel_name in names:
                rel_root = ET.fromstring(archive.read(rel_name))
                for rel in rel_root.iter(PR + "Relationship"):
                    target = rel.attrib.get("Target", "")
                    if "../media/" in target:
                        media.append("ppt/media/" + target.rsplit("/", 1)[-1])
            slides.append(
                {
                    "number": int(re.search(r"\d+", slide_name).group()),
                    "text": text,
                    "shapes": shapes,
                    "media": media,
                }
            )
        media_entries = [entry for entry in entries if entry.filename.startswith("ppt/media/")]
        largest = sorted(media_entries, key=lambda entry: entry.file_size, reverse=True)[:40]
        return {
            "slides": slides,
            "entry_count": len(entries),
            "media_count": len(media_entries),
            "media_uncompressed_bytes": sum(entry.file_size for entry in media_entries),
            "largest_media": [
                {
                    "name": entry.filename,
                    "bytes": entry.file_size,
                    "compressed_bytes": entry.compress_size,
                }
                for entry in largest
            ],
        }


def main() -> None:
    if len(sys.argv) != 6:
        raise SystemExit("Usage: script output-dir docx pdf xlsx pptx")
    output = Path(sys.argv[1])
    output.mkdir(parents=True, exist_ok=True)
    docx, pdf, xlsx, pptx = map(Path, sys.argv[2:])
    (output / "brand-philosophy.txt").write_text(docx_text(docx), encoding="utf-8")
    pdf_content, pdf_pages = pdf_text(pdf)
    (output / "brochure.txt").write_text(pdf_content, encoding="utf-8")
    (output / "brochure-images.json").write_text(
        json.dumps(pdf_image_manifest(pdf), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    extract_pdf_images(pdf, output / "brochure-images")
    (output / "specification.json").write_text(
        json.dumps(xlsx_text(xlsx), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    manifest = pptx_manifest(pptx)
    (output / "presentation.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "pdf_pages": pdf_pages,
                "pptx_slides": len(manifest["slides"]),
                "pptx_media": manifest["media_count"],
                "pptx_media_uncompressed_bytes": manifest["media_uncompressed_bytes"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
